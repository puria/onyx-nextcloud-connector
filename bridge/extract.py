"""Text extraction for supported formats: pdf, docx, txt, md."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ExtractionResult:
    text: str | None  # None = nothing extractable (e.g. scanned PDF)
    note: str | None


def extract_text(path: Path, suffix: str) -> ExtractionResult:
    suffix = suffix.lower()
    try:
        if suffix == ".pdf":
            return _pdf(path)
        if suffix == ".docx":
            return _docx(path)
        if suffix in {".txt", ".md"}:
            return _plaintext(path)
        return ExtractionResult(text=None, note="unsupported format")
    except Exception as exc:  # corrupt/unreadable file must not kill the run
        logger.warning("Extraction failed for %s: %s", path.name, type(exc).__name__)
        return ExtractionResult(text=None, note=f"extraction error: {type(exc).__name__}")


def _pdf(path: Path) -> ExtractionResult:
    from pypdf import PdfReader

    pages: list[str] = []
    with path.open("rb") as fh:
        reader = PdfReader(fh)
        for page in reader.pages:
            pages.append(page.extract_text() or "")
    text = "\n\n".join(p.strip() for p in pages if p and p.strip())
    if not text.strip():
        return ExtractionResult(text=None, note="scanned PDF without extractable text")
    return ExtractionResult(text=text, note=None)


def _docx(path: Path) -> ExtractionResult:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    text = "\n".join(parts)
    if not text.strip():
        return ExtractionResult(text=None, note="docx without extractable text")
    return ExtractionResult(text=text, note=None)


def _plaintext(path: Path) -> ExtractionResult:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if not text.strip():
        return ExtractionResult(text=None, note="empty file")
    return ExtractionResult(text=text, note=None)
