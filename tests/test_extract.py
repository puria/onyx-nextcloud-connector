"""Unit tests for text extraction."""

from __future__ import annotations

import io

from pypdf import PdfWriter

from bridge.extract import extract_text


def test_txt_extraction(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("hello text", encoding="utf-8")
    result = extract_text(path, ".txt")
    assert result.text == "hello text" and result.note is None


def test_markdown_extraction(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# Title\nbody", encoding="utf-8")
    assert extract_text(path, ".md").text.startswith("# Title")


def test_docx_extraction(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("first paragraph")
    document.add_paragraph("second paragraph")
    path = tmp_path / "a.docx"
    document.save(str(path))
    text = extract_text(path, ".docx").text
    assert "first paragraph" in text and "second paragraph" in text


def test_scanned_pdf_reports_no_text(tmp_path):
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(buffer)
    path = tmp_path / "scan.pdf"
    path.write_bytes(buffer.getvalue())
    result = extract_text(path, ".pdf")
    assert result.text is None
    assert result.note and "scanned PDF" in result.note


def test_corrupt_file_reports_error(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx")
    result = extract_text(path, ".docx")
    assert result.text is None and result.note is not None
